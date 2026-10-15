"""
QC & Fault Tracker Export Engine.
Generates in-memory CSV (encoded with UTF-8 BOM: utf-8-sig) and formatted Microsoft Excel (.xlsx)
spreadsheets with formula injection sanitization (CWE-1236) and severity color coding.
"""

import io
import csv
from typing import List, Dict, Any, Optional

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


def sanitize_csv_cell(val: Any) -> str:
    """
    Prevents CSV / Spreadsheet Formula Injection (CWE-1236).
    If a cell value begins with =, +, -, @, \\t, or \\r, prefix it with a single quote (').
    Also normalizes None values to empty strings.
    """
    if val is None:
        return ""
    text = str(val)
    if text and text[0] in ("=", "+", "-", "@", "\t", "\r"):
        return f"'{text}"
    return text


def build_qc_export_data(fault_logs: List[Dict[str, Any]]) -> tuple[List[str], List[List[str]]]:
    """
    Builds clean, sanitized header and row lists from fault log records.
    """
    headers = [
        "Fault ID",
        "Project Code",
        "Severity",
        "Status",
        "Labeler ID",
        "Labeler Name",
        "Team Name",
        "Team Leader ID",
        "Fault Type",
        "Description",
        "Screenshot URL",
        "Logged By ID",
        "Logged By Name",
        "Created At",
        "Acknowledged At"
    ]

    rows: List[List[str]] = []
    for f in fault_logs:
        row = [
            sanitize_csv_cell(f.get("id")),
            sanitize_csv_cell(f.get("project_code")),
            sanitize_csv_cell((f.get("severity") or "").upper()),
            sanitize_csv_cell((f.get("status") or "open").capitalize()),
            sanitize_csv_cell(f.get("labeler_id")),
            sanitize_csv_cell(f.get("labeler_name")),
            sanitize_csv_cell(f.get("team_name")),
            sanitize_csv_cell(f.get("leader_id") or "N/A"),
            sanitize_csv_cell(f.get("fault_type") or "General"),
            sanitize_csv_cell(f.get("description")),
            sanitize_csv_cell(f.get("image_url") or "None"),
            sanitize_csv_cell(f.get("logged_by_id")),
            sanitize_csv_cell(f.get("logged_by_name")),
            sanitize_csv_cell(f.get("created_at")),
            sanitize_csv_cell(f.get("acknowledged_at") or "Pending")
        ]
        rows.append(row)

    return headers, rows


def generate_qc_csv_bytes(fault_logs: List[Dict[str, Any]]) -> bytes:
    """
    Renders fault logs as CSV bytes with UTF-8 BOM encoding.
    """
    headers, rows = build_qc_export_data(fault_logs)
    output = io.StringIO()
    writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(row)

    csv_text = output.getvalue()
    return csv_text.encode("utf-8-sig")


def generate_qc_xlsx_bytes(
    fault_logs: List[Dict[str, Any]],
    sheet_title: str = "QC Fault Logs"
) -> bytes:
    """
    Renders fault logs as a Microsoft Excel workbook (.xlsx) with styled headers,
    borders, severity color-coded badges, and auto-adjusted column widths.
    """
    if not OPENPYXL_AVAILABLE:
        # Fallback to UTF-8 BOM CSV if openpyxl is not available
        return generate_qc_csv_bytes(fault_logs)

    headers, rows = build_qc_export_data(fault_logs)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]

    # Palette
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10, color="0F172A")

    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )

    # Severity fills
    minor_fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")     # soft green
    major_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")     # soft yellow
    critical_fill = PatternFill(start_color="FFE4E6", end_color="FFE4E6", fill_type="solid")  # soft red

    # Status fills
    open_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")      # light rose
    ack_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")       # light emerald

    # Write Headers
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[1].height = 28

    # Write Rows
    for row_idx, row_data in enumerate(rows, start=2):
        ws.append(row_data)
        ws.row_dimensions[row_idx].height = 20

        # Severity index is 2 (0-indexed) -> col 3
        # Status index is 3 (0-indexed) -> col 4
        sev_val = str(row_data[2]).upper()
        status_val = str(row_data[3]).lower()

        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

            # Apply severity styling
            if col_idx == 3:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if "CRITICAL" in sev_val:
                    cell.fill = critical_fill
                    cell.font = Font(name="Segoe UI", size=10, bold=True, color="991B1B")
                elif "MAJOR" in sev_val:
                    cell.fill = major_fill
                    cell.font = Font(name="Segoe UI", size=10, bold=True, color="92400E")
                elif "MINOR" in sev_val:
                    cell.fill = minor_fill
                    cell.font = Font(name="Segoe UI", size=10, bold=True, color="065F46")

            # Apply status styling
            if col_idx == 4:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if "acknowledged" in status_val:
                    cell.fill = ack_fill
                    cell.font = Font(name="Segoe UI", size=10, bold=True, color="166534")
                else:
                    cell.fill = open_fill
                    cell.font = Font(name="Segoe UI", size=10, bold=True, color="991B1B")

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 45)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
