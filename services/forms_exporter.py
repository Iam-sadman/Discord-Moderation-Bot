"""
Forms Export Engine.
Generates in-memory CSV (encoded with UTF-8 BOM: utf-8-sig) and formatted Microsoft Excel (.xlsx)
spreadsheets with dynamic question headers, formula injection sanitization, and review filtering.
"""

import io
import csv
import json
from typing import List, Dict, Any, Optional

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


def sanitize_csv_cell(val: Any) -> str:
    """
    Prevents CSV / Spreadsheet Formula Injection (CWE-1236).
    If a cell value begins with =, +, -, @, \\t, or \\r, prefix it with a single quote (').
    Also normalizes None values to empty strings.
    """
    if val is None:
        return ""
    text = str(val)
    if text and text[0] in ("=", "+", "-", "@", "\t", "\r"):
        return f"'{text}"
    return text


def build_export_data(
    form: Dict[str, Any],
    questions: List[Dict[str, Any]],
    submissions: List[Dict[str, Any]]
) -> tuple[List[str], List[List[str]]]:
    """
    Builds clean, sanitized header and row lists from form metadata, questions, and submissions.
    Dynamically adjusts columns depending on whether review_enabled is True or False.
    """
    review_enabled = bool(form.get("review_enabled", False))

    # Base metadata headers
    headers = ["Submission ID", "Status", "Submitted At", "User ID", "Username"]

    # Review headers only if review system is enabled for this form
    if review_enabled:
        headers.extend(["Reviewer ID", "Reviewer Name", "Review Note", "Reviewed At"])

    # Dynamic question headers: Q1: Label, Q2: Label...
    # Maintain question ID lookup for quick answer mapping
    question_map = {}
    for idx, q in enumerate(questions, start=1):
        q_label = q.get("label", f"Question {idx}").strip()
        headers.append(f"Q{idx}: {q_label}")
        question_map[str(q["id"])] = idx - 1  # 0-indexed offset within question columns

    rows: List[List[str]] = []
    for sub in submissions:
        sub_id = sanitize_csv_cell(sub.get("id"))
        status = sanitize_csv_cell(sub.get("status", "submitted").capitalize())
        created_at = sanitize_csv_cell(sub.get("created_at"))
        user_id = sanitize_csv_cell(sub.get("user_id"))
        username = sanitize_csv_cell(f"{sub.get('user_name', '')}#{sub.get('user_discriminator', '0')}")

        row = [sub_id, status, created_at, user_id, username]

        if review_enabled:
            rev_id = sanitize_csv_cell(sub.get("reviewer_id") or "")
            rev_name = sanitize_csv_cell(sub.get("reviewer_name") or "")
            rev_note = sanitize_csv_cell(sub.get("review_note") or "")
            rev_at = sanitize_csv_cell(sub.get("reviewed_at") or "")
            row.extend([rev_id, rev_name, rev_note, rev_at])

        # Parse answers
        answers_dict = sub.get("answers")
        if not isinstance(answers_dict, dict):
            try:
                answers_dict = json.loads(sub.get("answers_json") or "{}")
            except Exception:
                answers_dict = {}

        # Fill question responses in correct question position order
        q_answers = [""] * len(questions)
        for q_idx, q in enumerate(questions):
            qid_str = str(q["id"])
            q_label = q.get("label", "")
            # Support answer lookup by question id, string id, or label
            ans = answers_dict.get(qid_str) or answers_dict.get(q["id"]) or answers_dict.get(q_label) or ""
            q_answers[q_idx] = sanitize_csv_cell(ans)

        row.extend(q_answers)
        rows.append(row)

    return headers, rows


def generate_csv_bytes(
    form: Dict[str, Any],
    questions: List[Dict[str, Any]],
    submissions: List[Dict[str, Any]]
) -> io.BytesIO:
    """
    Generates an in-memory CSV file encoded with utf-8-sig (UTF-8 with BOM)
    for seamless opening in Microsoft Excel without corrupting emojis, Bengali, or Unicode characters.
    """
    headers, rows = build_export_data(form, questions, submissions)

    output = io.StringIO()
    writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")
    writer.writerow(headers)
    for r in rows:
        writer.writerow(r)

    # Convert to BytesIO with utf-8-sig
    byte_output = io.BytesIO(output.getvalue().encode("utf-8-sig"))
    byte_output.seek(0)
    return byte_output


def generate_xlsx_bytes(
    form: Dict[str, Any],
    questions: List[Dict[str, Any]],
    submissions: List[Dict[str, Any]]
) -> io.BytesIO:
    """
    Generates a professionally styled Microsoft Excel (.xlsx) workbook in-memory.
    Falls back to CSV bytes if openpyxl is not installed.
    """
    if not OPENPYXL_AVAILABLE:
        return generate_csv_bytes(form, questions, submissions)

    headers, rows = build_export_data(form, questions, submissions)

    wb = openpyxl.Workbook()
    ws = wb.active
    form_title = (form.get("title") or "Form Submissions")[:31]  # Excel tab name max 31 chars
    # Remove characters forbidden in Excel worksheet titles: \ / ? * : [ ]
    for ch in ("\\", "/", "?", "*", ":", "[", "]"):
        form_title = form_title.replace(ch, " ")
    ws.title = form_title.strip() or "Submissions"

    # Define professional styling tokens
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)
    border_side = Side(style="thin", color="E2E8F0")
    thin_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    data_alignment = Alignment(vertical="center", wrap_text=True)

    # Status color highlights
    status_fills = {
        "Approved": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
        "Rejected": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
        "Pending": PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
        "Submitted": PatternFill(start_color="E0E7FF", end_color="E0E7FF", fill_type="solid")
    }

    # Write Header row
    ws.append(headers)
    ws.row_dimensions[1].height = 28
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border

    # Write Data rows
    for row_idx, r in enumerate(rows, start=2):
        ws.append(r)
        ws.row_dimensions[row_idx].height = 22
        status_val = r[1] if len(r) > 1 else ""
        row_status_fill = status_fills.get(status_val)

        for col_num in range(1, len(r) + 1):
            cell = ws.cell(row=row_idx, column=col_num)
            cell.font = data_font
            cell.alignment = data_alignment
            cell.border = thin_border
            if col_num == 2 and row_status_fill:
                cell.fill = row_status_fill

    # Auto-fit column widths with sensible min/max constraints
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        # Bounded between 12 and 50 characters
        ws.column_dimensions[col_letter].width = max(12, min(max_len + 4, 50))

    byte_output = io.BytesIO()
    wb.save(byte_output)
    byte_output.seek(0)
    return byte_output
