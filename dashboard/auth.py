"""
Dashboard User Authentication & Excel Sheet Manager.
Manages user IDs, access keys, roles, and status using an Excel workbook (data/dashboard_users.xlsx).
"""

import os
import logging
from datetime import datetime
from typing import Optional, Dict, Any, List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

logger = logging.getLogger("DashboardAuth")

USERS_EXCEL_PATH = os.path.join(os.path.dirname(os.path.dirname(__file__)), "data", "dashboard_users.xlsx")


def init_users_sheet():
    """Initializes the Excel user registry if it does not already exist."""
    os.makedirs(os.path.dirname(USERS_EXCEL_PATH), exist_ok=True)
    if os.path.exists(USERS_EXCEL_PATH):
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Authorized Users"

    headers = ["User ID", "Access Key", "Full Name", "Role", "Status", "Created At", "Notes"]
    ws.append(headers)

    # Style header row
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    # Add default master admin account
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    default_users = [
        ["admin", "admin-key-2026", "Master Admin", "Administrator", "Active", now_str, "Default system administrator key"],
        ["supervisor", "staff-access-2026", "Team Supervisor", "Moderator", "Active", now_str, "Dashboard staff viewer key"]
    ]

    for row_data in default_users:
        ws.append(row_data)

    # Format ID and Key columns as Text to prevent scientific notation
    for col_idx in (1, 2):
        for row_idx in range(1, 100):
            ws.cell(row_idx, col_idx).number_format = '@'

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    wb.save(USERS_EXCEL_PATH)
    logger.info(f"Initialized dashboard users registry at {USERS_EXCEL_PATH}")


def _clean_excel_str(val) -> str:
    """Cleans Excel cell values, preventing float/scientific notation for IDs/Keys."""
    if val is None:
        return ""
    if isinstance(val, (int, float)):
        # If float with no decimal part or large integer
        try:
            int_val = int(val)
            if abs(float(int_val) - float(val)) < 1e-5:
                return str(int_val).strip()
            # If large float
            return f"{val:.0f}".strip()
        except (ValueError, OverflowError):
            pass
    return str(val).strip()


def load_all_users() -> List[Dict[str, Any]]:
    """Reads all user rows from the Excel workbook."""
    init_users_sheet()
    users = []
    try:
        wb = openpyxl.load_workbook(USERS_EXCEL_PATH, data_only=True)
        ws = wb.active

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            user_id = _clean_excel_str(row[0])
            access_key = _clean_excel_str(row[1]) if len(row) > 1 and row[1] is not None else ""
            name = str(row[2]).strip() if len(row) > 2 and row[2] is not None else user_id
            role = str(row[3]).strip() if len(row) > 3 and row[3] is not None else "Viewer"
            status = str(row[4]).strip() if len(row) > 4 and row[4] is not None else "Active"
            created_at = str(row[5]) if len(row) > 5 and row[5] is not None else ""
            notes = str(row[6]) if len(row) > 6 and row[6] is not None else ""

            users.append({
                "user_id": user_id,
                "access_key": access_key,
                "raw_access_key": str(row[1]).strip() if len(row) > 1 and row[1] is not None else "",
                "name": name,
                "role": role,
                "status": status,
                "created_at": created_at,
                "notes": notes
            })
    except Exception as e:
        logger.error(f"Error reading dashboard users Excel file: {e}")
    return users


def authenticate_user(user_id: str, access_key: str) -> Optional[Dict[str, Any]]:
    """
    Validates user credentials against the Excel registry.
    Returns user dictionary if valid and active, otherwise None.
    """
    if not user_id or not access_key:
        return None

    clean_id = user_id.strip()
    clean_key = access_key.strip()

    users = load_all_users()
    for u in users:
        if u["user_id"].lower() == clean_id.lower():
            key_matches = (
                u["access_key"] == clean_key or
                u.get("raw_access_key") == clean_key or
                (clean_key.isdigit() and u["access_key"].isdigit() and clean_key[:14] == u["access_key"][:14])
            )
            if key_matches:
                if u["status"].lower() == "active":
                    return u
    return None
